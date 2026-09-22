"""QA Excel exporter 的长期契约测试。"""

import io

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase

from core.models import Project, QADomainResult, QAReference
from core.quality.exporters.excel import export_qa_excel


User = get_user_model()


class QaExcelExporterTests(TestCase):
    def test_exporter_returns_expected_workbook_without_http_request(self):
        user = User.objects.create_user('qa-export-user')
        project = Project.objects.create(name='QA export project', owner=user)
        QAReference.objects.create(
            project=project,
            title='Exported Study',
            quality_method='QUADAS2',
            review_status='confirmed',
        )

        filename, content = export_qa_excel(project, 'QUADAS2')

        self.assertTrue(filename.endswith('.xlsx'))
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ['汇总统计', '评价明细', '证据记录', '多模型校验记录'],
        )
        self.assertEqual(workbook['汇总统计']['A2'].value, 'Exported Study')
        workbook.close()

    def test_summary_columns_follow_selected_method_configuration(self):
        user = User.objects.create_user('qa-nos-export-user')
        project = Project.objects.create(name='NOS export project', owner=user)
        ref = QAReference.objects.create(
            project=project,
            title='NOS Study',
            quality_method='NOS',
            review_status='confirmed',
        )
        QADomainResult.objects.create(
            qa_ref=ref,
            domain='selection',
            domain_name='Selection',
            bias_risk_result='high',
            applicability_result='na',
        )
        QADomainResult.objects.create(
            qa_ref=ref,
            domain='comparability',
            domain_name='Comparability',
            bias_risk_result='unclear',
            applicability_result='na',
        )
        QADomainResult.objects.create(
            qa_ref=ref,
            domain='outcome',
            domain_name='Outcome',
            bias_risk_result='low',
            applicability_result='na',
        )

        _, content = export_qa_excel(project, 'NOS')
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        rows = list(workbook['汇总统计'].iter_rows(values_only=True))
        workbook.close()

        self.assertEqual(
            rows[0],
            (
                '文献标题', '第一作者', '年份',
                '选择（Selection）_偏倚', '可比性（Comparability）_偏倚',
                '结局（Outcome）_偏倚', '整体审阅状态',
            ),
        )
        self.assertEqual(rows[1][3:6], ('high', 'unclear', 'low'))

    def test_unconfirmed_references_are_exported_only_when_requested(self):
        user = User.objects.create_user('qa-export-scope-user')
        project = Project.objects.create(name='QA export scope project', owner=user)
        QAReference.objects.create(
            project=project,
            title='Confirmed Study',
            quality_method='QUADAS2',
            review_status='confirmed',
        )
        QAReference.objects.create(
            project=project,
            title='Pending Study',
            quality_method='QUADAS2',
            review_status='partial',
        )

        _, confirmed_content = export_qa_excel(project, 'QUADAS2')
        _, all_content = export_qa_excel(project, 'QUADAS2', include_unconfirmed=True)

        confirmed_book = openpyxl.load_workbook(io.BytesIO(confirmed_content), read_only=True)
        all_book = openpyxl.load_workbook(io.BytesIO(all_content), read_only=True)
        confirmed_titles = [
            row[0] for row in list(confirmed_book['汇总统计'].iter_rows(values_only=True))[1:]
        ]
        all_titles = [
            row[0] for row in list(all_book['汇总统计'].iter_rows(values_only=True))[1:]
        ]
        confirmed_book.close()
        all_book.close()

        self.assertEqual(confirmed_titles, ['Confirmed Study'])
        self.assertEqual(all_titles, ['Confirmed Study', 'Pending Study'])
