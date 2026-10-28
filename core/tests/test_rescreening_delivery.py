"""Stages 8–10: real file exports and incremental QA handoff, without provider calls."""
import json
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Project, QAReference, Task
from core.fulltext.access import asset_for_qa
from core.fulltext.services import asset_is_referenced
from core.rescreening.errors import RescreenError
from core.rescreening.models import RescreenExport, RescreenExportItem, QASourceBibliographySnapshot
from core.rescreening.services.exports import start_export, process_export, fingerprint, finish_task
from core.rescreening.services.qa_handoff import preview, handoff
from core.rescreening.services.workspace import save_review, reference_query, workspace_state, save_configuration
from core.rescreening.services.cleanup import cleanup_failed_imports
from core.tests.test_shared_fulltext import SharedFulltextFixture
from core.services.storage_quota_service import user_storage_usage_bytes


class RescreenDeliveryFixture(SharedFulltextFixture):
    def setUp(self):
        super().setUp()
        self.asset, refs = self.ready(references=True)
        self.ref = refs[0]
        self.ref.refresh_from_db()

    def review(self, **changes):
        row = reference_query(self.project).get(pk=self.ref.pk)
        data = dict(project=self.project, user=self.user, reference_id=row.pk, input_version=row.input_version,
            expected_version=row.current_review.version if row.current_review_id else 0,
            criteria_revision_id=workspace_state(self.project,self.user)['configuration']['criteria_revision_id'],
            suggestion_id=row.suggestion_id, suggestion_revision=row.suggestion_revision or 0,
            decision='included')
        return save_review(**(data | changes))

    def start(self, **changes):
        data = dict(project=self.project, user=self.user, format='xlsx', scope='all', mode='final',
                    idempotency_key=str(uuid4()), expected_fingerprint=fingerprint(self.project))
        with patch('core.rescreening.services.exports._dispatch'):
            return start_export(**(data | changes))

    def generate(self, **changes):
        job = self.start(**changes)
        process_export(job.pk, expected_task_id=job.task_id)
        Task.objects.filter(pk=job.task_id).update(status='completed')
        job.refresh_from_db()
        return job

    def transfer(self, **changes):
        data = dict(project=self.project, user=self.user, confirmation_token=preview(self.project)['confirmation_token'])
        return handoff(**(data | changes))


class RescreenExportTests(RescreenDeliveryFixture):
    def test_explicit_skip_is_durable_idempotent_and_requires_final_reviews(self):
        from core.rescreening.services.exports import skip_export
        data = dict(project=self.project, user=self.user, idempotency_key='skip', expected_fingerprint=fingerprint(self.project))
        with self.assertRaisesMessage(RescreenError, '最终确认'):
            skip_export(**data)
        self.review()
        data['expected_fingerprint'] = fingerprint(self.project)
        job = skip_export(**data)
        self.assertEqual(skip_export(**data).pk, job.pk)
        self.assertIsNone(job.task_id)
        state = workspace_state(self.project, self.user)
        self.assertTrue(state['steps'][5]['completed'])
        self.assertEqual(state['export_completion_kind'], 'skipped')
        self.review(note='更改后需重新确认跳过')
        self.assertFalse(workspace_state(self.project,self.user)['steps'][5]['completed'])

    def test_pending_only_draft_and_batch_range_are_explicit(self):
        job = self.generate(format='metadata', scope='pending', mode='draft', offset=0, limit=1)
        self.assertIn('_part_1_1.', job.download_name)
        with job.output.open('r') as output:
            record = json.loads(output.readline())
        self.assertEqual(record['decision'], 'pending')
        self.assertEqual(record['fulltext']['sha256'], self.asset.sha256)
        self.assertIn('field_mapping', record)
        self.assertIn('criteria', record)
        with self.assertRaisesMessage(RescreenError, '没有可导出'):
            self.start(mode='draft', offset=1, limit=1)

    def test_abandoned_private_spool_is_cleaned_without_touching_live_or_unrelated_folders(self):
        import os
        self.review()
        live = self.start()
        root = self.root / 'exports' / '.tmp'
        root.mkdir(parents=True)
        abandoned, active, unrelated = root/'job_999999_abandoned', root/f'job_{live.pk}_active', root/'unrelated'
        for folder in (abandoned, active, unrelated):
            folder.mkdir()
            os.utime(folder, (1,1))
        cleanup_failed_imports()
        self.assertFalse(abandoned.exists())
        self.assertTrue(active.exists())
        self.assertTrue(unrelated.exists())

    def test_malformed_export_options_are_400_not_internal_errors(self):
        response = self.client.post('/api/rescreening/exports/', {'project_id':self.project.pk,
            'format':['xlsx'], 'fingerprint':fingerprint(self.project)}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(RescreenExport.objects.exists())
        response = self.client.post('/api/rescreening/exports/', {'project_id':self.project.pk,
            'format':'xlsx', 'offset':10**99, 'fingerprint':fingerprint(self.project)}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_output_quota_failure_never_publishes_a_file(self):
        from core.services.storage_quota_service import StorageQuotaExceeded, StorageQuotaSnapshot
        self.review()
        job = self.start()
        with patch('core.rescreening.services.exports.enforce_user_storage_quota',
                   side_effect=StorageQuotaExceeded(StorageQuotaSnapshot(0, 0), 1)):
            with self.assertRaises(StorageQuotaExceeded):
                process_export(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.error_code, 'storage_quota_exceeded')
        self.assertFalse(job.output)
        self.assertFalse(workspace_state(self.project, self.user)['steps'][5]['completed'])

    def test_fingerprint_uses_aggregates_not_per_reference_cold_reads(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        from core.rescreening.models import RescreenReference
        fields = {field.name: getattr(self.ref, field.attname) for field in RescreenReference._meta.fields
                  if field.name not in ('id','source_key','created_at','updated_at')}
        # FK values are IDs for portable bulk_create, not extra per-row queries.
        for field in RescreenReference._meta.fields:
            if field.is_relation and field.name in fields:
                fields[field.attname] = fields.pop(field.name)
        RescreenReference.objects.bulk_create([RescreenReference(**fields, source_key=f'large:{i}') for i in range(1000)])
        with CaptureQueriesContext(connection) as queries:
            value = fingerprint(self.project)
        self.assertEqual(len(value), 64)
        self.assertLessEqual(len(queries), 8)
        self.assertFalse(any('plat_rescreen_reference_raw' in row['sql'] for row in queries))

    def test_cleanup_api_preserves_receipt_and_releases_logical_snapshot_quota(self):
        self.review()
        job = self.generate()
        before = user_storage_usage_bytes(self.user)
        size = job.output_size + job.snapshot_size
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(f'/api/rescreening/exports/{job.pk}/purge/', {}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(user_storage_usage_bytes(self.user), before-size)
        self.assertTrue(workspace_state(self.project, self.user)['steps'][5]['completed'])

    def test_public_project_stage_uses_current_completion_not_stale_db_status(self):
        from core.serializers import ProjectStageSerializer
        stage = self.project.stages.get(stage_key='SCREEN_2')
        self.review()
        self.generate()
        stage.refresh_from_db()
        self.assertEqual(ProjectStageSerializer(stage).data['status'], 'completed')
        self.assertEqual(ProjectStageSerializer(stage).data['progress_percentage'], 100.0)
        self.review(note='新版本')
        stage.refresh_from_db()
        self.assertEqual(ProjectStageSerializer(stage).data['status'], 'in_progress')

    def test_actual_worker_handler_completes_without_model_slots_or_fees(self):
        from core.executors.celery_tasks import execute_async_step
        self.review()
        job = self.start()
        with patch('core.services.concurrency_service.try_acquire') as acquire:
            self.assertTrue(execute_async_step.run(job.task_id, 'rescreen_export', self.project.pk))
        acquire.assert_not_called()
        job.refresh_from_db()
        job.task.refresh_from_db()
        self.assertEqual((job.status, job.task.status), ('completed', 'completed'))
        self.assertTrue(workspace_state(self.project,self.user)['steps'][5]['completed'])

    def test_final_requires_human_confirmation_and_draft_does_not_complete(self):
        with self.assertRaisesMessage(RescreenError, '只能导出草稿'):
            self.start()
        job = self.generate(mode='draft')
        self.assertEqual(job.status, 'completed')
        self.assertFalse(workspace_state(self.project,self.user)['steps'][5]['completed'])
        book = load_workbook(job.output.path, read_only=True)
        self.assertEqual(list(book['复筛结果'].values)[1][24], 'pending')
        book.close()

    def test_final_completion_survives_refresh_and_changes_invalidate_it(self):
        self.review()
        job = self.generate()
        self.assertTrue(workspace_state(self.project,self.user)['steps'][5]['completed'])
        self.review(decision='excluded', reason='年龄不符')
        state = workspace_state(self.project,self.user)
        self.assertFalse(state['steps'][5]['completed'])
        self.assertNotEqual(state['fingerprint'], job.fingerprint)
        data = self.client.get('/api/rescreening/exports/', {'project_id': self.project.pk}).json()['data']
        self.assertFalse(data['items'][0]['current'])
        self.assertIsNotNone(data['items'][0]['download_url'])

    def test_snapshot_is_frozen_even_when_review_changes_before_worker(self):
        self.review(note='冻结备注')
        job = self.start(format='metadata')
        self.review(decision='excluded', reason='新理由')
        process_export(job.pk)
        job.refresh_from_db()
        with job.output.open('r') as output:
            payload = json.loads(output.readline())
        self.assertEqual(payload['decision'], 'included')
        self.assertEqual(payload['review']['note'], '冻结备注')

    def test_all_formats_preserve_bibliography_and_raw_vendor_fields(self):
        self.ref.title = '=HYPERLINK("https://invalid.example","title")'
        self.ref.organization = 'University; department'
        self.ref.address = 'Beijing & London'
        self.ref.isbn = '978-1-2'
        self.ref.pmid = '12345'
        self.ref.keywords = ['甲', '乙']
        self.ref.save()
        # Simulate a parser's original vendor field, not a public mutation API.
        self.ref.raw_metadata_versions.update(raw_fields={'vendor_custom': {'a': ['extra']}})
        self.review(note='保持所有字段')
        for format in ('xlsx', 'xml', 'ris', 'metadata'):
            with self.subTest(format=format):
                job = self.generate(format=format)
                if format == 'xlsx':
                    book = load_workbook(job.output.path, read_only=True)
                    cells = list(book['复筛结果'].rows)
                    title_index = list(c.value for c in cells[0]).index('title')
                    self.assertEqual(cells[1][title_index].value, self.ref.title)
                    self.assertEqual(cells[1][title_index].data_type, 's')
                    raw = list(book['长字段与原始元数据'].values)[1][3]
                    self.assertIn('vendor_custom', raw)
                    book.close()
                else:
                    with job.output.open('r') as output:
                        text = output.read()
                    self.assertIn('vendor_custom', text)
                    self.assertIn('978-1-2', text)
                    if format == 'xml':
                        root = ET.fromstring(text)
                        self.assertIn('Beijing & London', root.find('Reference/bibliography').text)
                        from core.screening import parsers
                        parsed = list(parsers.iter_file(job.output.path))[0]
                        self.assertEqual(parsed['title'], self.ref.title)
                        self.assertEqual(parsed['isbn'], self.ref.isbn)
                    if format == 'ris':
                        self.assertIn('TI  - ' + self.ref.title, text)
                        self.assertLess(text.index('N1  - MetaFlow'), text.index('ER  -'))

    def test_excel_long_fields_are_split_not_truncated(self):
        self.ref.abstract = '中文 abstract ' * 6000
        self.ref.save(update_fields=['abstract'])
        self.review()
        job = self.generate()
        book = load_workbook(job.output.path, read_only=True)
        parts = [r for r in list(book['长字段与原始元数据'].values)[1:] if r[1] == 'abstract']
        self.assertEqual(''.join(r[3] for r in sorted(parts,key=lambda r:r[2])), self.ref.abstract)
        book.close()

    def test_pdf_zip_contains_safe_ids_manifest_and_checksums(self):
        self.review()
        job = self.generate(format='pdf')
        with ZipFile(job.output.path) as archive:
            manifest = json.loads(archive.read('manifest.json'))
            name = manifest['files'][0]['file']
            self.assertEqual(name, f'pdf/reference_{self.ref.pk}.pdf')
            self.assertTrue(archive.read(name).startswith(b'%PDF'))
            self.assertEqual(manifest['files'][0]['sha256'], self.asset.sha256)
        self.assertFalse(workspace_state(self.project,self.user)['steps'][5]['completed'])

    def test_size_bound_and_stop_do_not_publish_partial_files(self):
        self.review()
        with override_settings(RESCREEN_EXPORT_MAX_BYTES=5):
            with self.assertRaisesMessage(RescreenError, '超过'):
                self.start(format='pdf')
        job = self.start(format='metadata')
        with override_settings(RESCREEN_EXPORT_MAX_BYTES=5):
            with self.assertRaisesMessage(RescreenError, '超过'):
                process_export(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertFalse(job.output)
        Task.objects.filter(pk=job.task_id).update(status='failed')
        job = self.start()
        with self.assertRaisesMessage(RescreenError, '停止'):
            process_export(job.pk, check_stop=lambda: True)
        job.refresh_from_db()
        self.assertEqual(job.status, 'cancelled')
        self.assertFalse(list((self.root/'exports').rglob('*.xlsx')))

    def test_idempotency_and_concurrent_start_rejection(self):
        self.review()
        job = self.start(idempotency_key='same')
        self.assertEqual(self.start(idempotency_key='same').pk, job.pk)
        with self.assertRaisesMessage(RescreenError, '幂等键'):
            self.start(idempotency_key='same', format='xml')
        with self.assertRaisesMessage(RescreenError, '已有导出'):
            self.start()

    def test_stale_source_and_stale_worker_are_rejected(self):
        self.review()
        with self.assertRaisesMessage(RescreenError, '结果已变化'):
            self.start(expected_fingerprint='old')
        job = self.start()
        with self.assertRaisesMessage(RescreenError, '任务已变化'):
            process_export(job.pk, expected_task_id=99999)

    def test_exports_are_private_paginated_and_counted_in_quota(self):
        self.review()
        before = user_storage_usage_bytes(self.user)
        job = self.generate()
        self.assertEqual(user_storage_usage_bytes(self.user)-before, job.output_size+job.snapshot_size)
        with self.assertRaises(ValueError):
            job.output.url
        response = self.client.get(f'/api/rescreening/exports/{job.pk}/download/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        response.close()
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/rescreening/exports/{job.pk}/download/').status_code, 404)
        self.assertEqual(self.client.get('/api/rescreening/exports/', {'project_id':self.project.pk}).status_code, 403)

    def test_retention_removes_bytes_and_details_but_not_success_receipt(self):
        self.review()
        job = self.generate()
        path = Path(job.output.path)
        with self.captureOnCommitCallbacks(execute=True):
            cleanup_failed_imports(now=timezone.now()+timedelta(days=8))
        job.refresh_from_db()
        self.assertFalse(path.exists())
        self.assertFalse(job.items.exists())
        self.assertEqual(job.output_size, 0)
        self.assertTrue(workspace_state(self.project,self.user)['steps'][5]['completed'])
        self.assertEqual(self.client.get(f'/api/rescreening/exports/{job.pk}/download/').status_code, 410)

    def test_task_terminal_sync_and_project_cascade_remove_private_output(self):
        self.review()
        job = self.start()
        finish_task(job.task, cancelled=True)
        job.refresh_from_db()
        self.assertEqual(job.status, 'cancelled')
        Task.objects.filter(pk=job.task_id).update(status='stopped')
        job = self.generate()
        output = Path(job.output.path)
        with self.captureOnCommitCallbacks(execute=True):
            self.project.delete()
        self.assertFalse(output.exists())
        self.assertFalse(RescreenExportItem.objects.exists())


class RescreenQAHandoffTests(RescreenDeliveryFixture):
    def test_qa_source_appendix_over_budget_returns_a_clear_413(self):
        self.review()
        self.transfer()
        qa = QAReference.objects.get()
        qa.quality_method, qa.review_status = 'QUADAS2', 'confirmed'
        qa.save()
        self.client.force_login(self.user)
        with override_settings(RESCREEN_EXPORT_MAX_BYTES=1):
            response = self.client.post('/api/qa/export/excel/', {'project_id':self.project.pk, 'quality_method':'QUADAS2'}, format='json')
        self.assertEqual(response.status_code, 413)
        self.assertIn('来源元数据', response.json()['error'])

    def test_qa_excel_keeps_full_source_fields_and_original_vendor_metadata(self):
        from io import BytesIO
        from core.quality.exporters.excel import export_qa_excel
        self.ref.organization = '完整来源单位'
        self.ref.title = '=1+1'
        self.ref.save(update_fields=['organization', 'title'])
        self.review()
        self.transfer()
        qa = QAReference.objects.get()
        qa.quality_method, qa.review_status = 'QUADAS2', 'confirmed'
        qa.save()
        filename, content = export_qa_excel(self.project, 'QUADAS2')
        book = load_workbook(BytesIO(content), read_only=True)
        rows = list(book['复筛来源元数据'].values)[1:]
        source = json.loads(''.join(row[2] for row in rows))
        self.assertEqual(source['bibliography']['organization'], '完整来源单位')
        self.assertIn('raw_fields', source['raw_metadata'])
        self.assertIn('review', source['provenance'])
        self.assertEqual(book['汇总统计']['A2'].value, '=1+1')
        self.assertEqual(book['汇总统计']['A2'].data_type, 's')
        book.close()

    def test_preview_rejects_a_source_that_changes_during_scan(self):
        with patch('core.rescreening.services.qa_handoff.fingerprint', side_effect=['before','after']):
            with self.assertRaisesMessage(RescreenError, '预览期间变化'):
                preview(self.project)

    def test_handoff_feeds_existing_qa_evidence_pipeline_and_authenticated_download(self):
        from core.quality.services.evidence_retrieval import build_reference_evidence_packages
        self.review()
        self.transfer()
        qa = QAReference.objects.get()
        qa.quality_method = 'QUADAS2'
        qa.save(update_fields=['quality_method'])
        packages = build_reference_evidence_packages(qa)
        self.assertTrue(packages)
        self.client.force_login(self.user)
        response = self.client.get('/api/qa/refs/', {'project_id':self.project.pk, 'page':1})
        self.assertEqual(response.status_code, 200)
        row = response.json()['data']['results'][0]
        self.assertEqual(row['fulltext_asset']['kind'], 'shared')
        download = self.client.get(row['fulltext_url'])
        self.assertEqual(download.status_code, 200)
        download.close()

    def test_only_final_included_imports_with_complete_snapshot_and_shared_pdf(self):
        self.assertEqual(preview(self.project)['counts']['new'], 0)
        self.review()
        size = user_storage_usage_bytes(self.user)
        result = self.transfer()
        qa = QAReference.objects.get()
        self.assertEqual(result['imported'], 1)
        self.assertEqual(qa.source_rescreen_reference_id, self.ref.pk)
        self.assertEqual(asset_for_qa(qa).pk, self.asset.pk)
        self.assertEqual(qa.fulltext_status, 'available')
        snapshot = qa.rescreen_source_snapshot
        self.assertEqual(snapshot.bibliography['title'], self.ref.title)
        self.assertIn('raw_fields', snapshot.raw_metadata)
        self.assertEqual(snapshot.provenance['review']['decision'], 'included')
        self.assertEqual(user_storage_usage_bytes(self.user), size)
        self.assertTrue(asset_is_referenced(self.asset))

    def test_repeated_import_preserves_evaluation_method_and_review(self):
        self.review()
        self.transfer()
        qa = QAReference.objects.get()
        qa.quality_method, qa.ai_eval_status, qa.review_status = 'QUADAS2', 'completed', 'confirmed'
        qa.save()
        self.assertEqual(self.transfer()['imported'], 0)
        qa.refresh_from_db()
        self.assertEqual((qa.quality_method, qa.ai_eval_status, qa.review_status), ('QUADAS2', 'completed', 'confirmed'))
        self.assertEqual(QAReference.objects.count(), 1)

    def test_source_changes_require_explicit_ack_and_never_overwrite_snapshot(self):
        self.review(note='原确认')
        self.transfer()
        self.review(note='新确认')
        info = preview(self.project)
        self.assertEqual(info['counts']['changed'], 1)
        with self.assertRaisesMessage(RescreenError, '确认来源差异'):
            self.transfer()
        self.transfer(acknowledge_differences=True)
        self.assertEqual(QASourceBibliographySnapshot.objects.get().provenance['review']['note'], '原确认')
        self.review(decision='excluded', reason='后续排除')
        self.assertEqual(preview(self.project)['counts']['no_longer_included'], 1)
        self.transfer(acknowledge_differences=True)
        self.assertEqual(QAReference.objects.count(), 1)

    def test_stale_confirmation_does_not_import_and_existing_legacy_qa_survives(self):
        legacy = QAReference.objects.create(project=self.project, title='旧 QA', quality_method='NOS')
        self.review()
        token = preview(self.project)['confirmation_token']
        self.review(note='并发修改')
        with self.assertRaisesMessage(RescreenError, '重新预览'):
            self.transfer(confirmation_token=token)
        self.transfer()
        legacy.refresh_from_db()
        self.assertEqual(legacy.quality_method, 'NOS')
        self.assertEqual(QAReference.objects.count(), 2)

    def test_unready_pdf_does_not_block_manual_included_or_invent_fulltext(self):
        self.ref.current_binding = None
        self.ref.save(update_fields=['current_binding'])
        self.review()
        result = self.transfer()
        qa = QAReference.objects.get()
        self.assertEqual(result['no_fulltext'], 1)
        self.assertIsNone(qa.shared_fulltext_asset_id)
        self.assertEqual(qa.fulltext_status, 'pending')

    def test_qa_busy_blocks_handoff_and_permission_is_project_scoped(self):
        self.review()
        task = Task.objects.create(project=self.project, task_type='qa_eval', status='running', created_by=self.user)
        with self.assertRaisesMessage(RescreenError, '已有活动任务'):
            self.transfer()
        task.delete()
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get('/api/rescreening/qa-handoff/', {'project_id':self.project.pk}).status_code, 403)

    def test_legacy_qa_import_endpoint_uses_new_non_destructive_screen2_path(self):
        self.review()
        self.client.force_login(self.user)
        response = self.client.post('/api/qa/refs/import/', {'project_id':self.project.pk, 'source_stage':'SCREEN_2',
            'confirmation_token':preview(self.project)['confirmation_token']}, format='json')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(QAReference.objects.count(), 1)
        self.assertTrue(QASourceBibliographySnapshot.objects.exists())
